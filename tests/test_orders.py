"""Order numbering, rendering, storage and cancellation. docs/05_TESTING.md §5."""

import re
from datetime import date, timedelta

import pytest
from django.test import override_settings
from django.utils import timezone
from docx import Document

from assignments.models import Assignment, AssignmentRole, AssignmentStatus
from audit.models import AuditLogEntry
from core.exceptions import (
    CancellationReasonRequired,
    NotPrincipal,
    OrderAlreadyCancelled,
    OrderDateInvalid,
    OrderNotCancellable,
    PreviewOutdated,
    TemplateMissing,
)
from orders.models import Order, OrderSequence, OrderType
from orders.services import (
    SPECIMEN_LINE,
    cancel_order,
    issue_assignment_order,
    issue_relinquishment_order,
    peek_next_order_no,
    preview_assignment_order,
)

pytestmark = pytest.mark.django_db


def document_text(order):
    with order.generated_file.open("rb") as handle:
        document = Document(handle)
    parts = [p.text for p in document.paragraphs]
    parts += [c.text for t in document.tables for row in t.rows for c in row.cells]
    return "\n".join(parts)


class TestNumbering:
    def test_format_matches_the_specified_pattern(self, faculty, committee, issue_order):
        order = issue_order(faculty=faculty, committee=committee)
        assert re.fullmatch(r"GPB/TPU/2026-27/\d{3}", order.order_no)

    def test_numbers_run_in_sequence(
        self, faculty, other_faculty, third_faculty, committee, issue_order
    ):
        numbers = [
            issue_order(
                faculty=person,
                committee=committee,
                role=AssignmentRole.MEMBER,
                order_date=date(2026, 7, day),
            ).order_no
            for day, person in enumerate([faculty, other_faculty, third_faculty], start=1)
        ]
        assert [n.rsplit("/", 1)[1] for n in numbers] == ["001", "002", "003"]

    def test_each_committee_has_its_own_sequence(
        self, faculty, committee, exam_committee, issue_order
    ):
        first = issue_order(faculty=faculty, committee=committee)
        second = issue_order(faculty=faculty, committee=exam_committee)
        assert first.order_no.endswith("001")
        assert second.order_no.endswith("001")

    def test_each_academic_year_has_its_own_sequence(
        self, faculty, committee, issue_order, last_year
    ):
        old = issue_order(
            faculty=faculty,
            committee=committee,
            order_date=date(2025, 7, 1),
            academic_year=last_year,
        )
        assert "2025-26/001" in old.order_no

    def test_a_cancelled_number_is_never_reused(
        self, principal, faculty, other_faculty, committee, issue_order
    ):
        first = issue_order(faculty=faculty, committee=committee)
        second = issue_order(
            faculty=other_faculty,
            committee=committee,
            role=AssignmentRole.MEMBER,
            order_date=date(2026, 7, 2),
        )
        cancel_order(order=second, reason="Issued in error.", issued_by=principal)
        third = issue_order(
            faculty=other_faculty,
            committee=committee,
            role=AssignmentRole.MEMBER,
            order_date=date(2026, 7, 5),
        )
        assert first.order_no.endswith("001")
        assert second.order_no.endswith("002")
        assert third.order_no.endswith("004")  # 003 was the corrigendum

    def test_peek_does_not_consume_a_number(self, committee, year):
        assert peek_next_order_no(year, committee).endswith("001")
        assert peek_next_order_no(year, committee).endswith("001")
        assert not OrderSequence.objects.exists()


class TestRendering:
    def test_the_document_contains_the_order_details(
        self, faculty, committee, issue_order
    ):
        order = issue_order(faculty=faculty, committee=committee)
        text = document_text(order)
        assert faculty.full_name in text
        assert faculty.salutation in text
        assert committee.name in text
        assert order.order_no in text
        assert order.order_date.strftime("%d-%m-%Y") in text
        assert "Presiding Officer" in text

    def test_no_placeholder_survives_rendering(self, faculty, committee, issue_order):
        assert "{{" not in document_text(issue_order(faculty=faculty, committee=committee))

    def test_a_template_using_optional_placeholders_still_renders(
        self, faculty, exam_committee, issue_order
    ):
        order = issue_order(
            faculty=faculty, committee=exam_committee, remarks="Duties to follow."
        )
        assert "Duties to follow." in document_text(order)

    def test_a_member_table_prints_every_person_appointed(
        self, faculty, other_faculty, third_faculty, external_member, table_committee, issue_order
    ):
        """The college's real orders name several people in one table."""
        order = issue_order(
            committee=table_committee,
            appointments=[
                (faculty, AssignmentRole.PRESIDING_OFFICER),
                (other_faculty, AssignmentRole.MEMBER),
                (third_faculty, AssignmentRole.MEMBER),
                (external_member, AssignmentRole.EXTERNAL_MEMBER),
            ],
        )
        with order.generated_file.open("rb") as handle:
            document = Document(handle)
        rows = document.tables[0].rows
        printed = [[cell.text.strip() for cell in row.cells] for row in rows[1:]]

        assert len(printed) == 4, "one row per person"
        assert [row[0] for row in printed] == ["1", "2", "3", "4"]
        assert printed[0][1].endswith(faculty.full_name)
        assert printed[0][3] == "Presiding Officer"
        assert printed[3][1].endswith(external_member.full_name)
        assert printed[3][2] == external_member.organisation
        assert printed[3][3] == "External Member"
        assert "{{" not in document_text(order)

    def test_a_reference_block_is_printed_when_given(
        self, faculty, table_committee, issue_order
    ):
        order = issue_order(
            committee=table_committee,
            faculty=faculty,
            references="GTU circular 2026-27/456 dated 04-02-2026",
        )
        assert "GTU circular 2026-27/456" in document_text(order)

    def test_the_reference_line_is_absent_when_not_given(
        self, faculty, table_committee, issue_order
    ):
        order = issue_order(committee=table_committee, faculty=faculty)
        assert "Reference:" not in document_text(order)

    def test_the_render_context_is_stored(self, faculty, committee, issue_order):
        order = issue_order(faculty=faculty, committee=committee)
        assert order.render_context["faculty_name"] == faculty.full_name
        assert order.render_context["order_no"] == order.order_no

    @override_settings(DEMO_MODE=True)
    def test_demo_mode_stamps_the_specimen_line_first(
        self, faculty, committee, issue_order
    ):
        order = issue_order(faculty=faculty, committee=committee)
        with order.generated_file.open("rb") as handle:
            document = Document(handle)
        assert document.paragraphs[0].text == SPECIMEN_LINE

    @override_settings(DEMO_MODE=False)
    def test_without_demo_mode_there_is_no_specimen_line(
        self, faculty, committee, issue_order
    ):
        order = issue_order(faculty=faculty, committee=committee)
        assert SPECIMEN_LINE not in document_text(order)

    def test_nothing_can_be_issued_without_an_installed_template(
        self, faculty, committee, issue_order, order_template
    ):
        from orders.models import OrderTemplate

        OrderTemplate.objects.update(is_active=False)
        with pytest.raises(TemplateMissing):
            issue_order(faculty=faculty, committee=committee)
        assert Order.objects.count() == 0

    def test_every_committee_uses_the_same_template(
        self, faculty, committee, exam_committee, issue_order, order_template
    ):
        first = issue_order(faculty=faculty, committee=committee)
        second = issue_order(faculty=faculty, committee=exam_committee)
        assert first.template_used == second.template_used == order_template

    def test_a_new_committee_can_issue_at_once(self, faculty, issue_order, db):
        from committees.models import Committee

        fresh = Committee.objects.create(code="NEW", name="Brand New Committee")
        assert issue_order(faculty=faculty, committee=fresh).generated_file

    def test_a_stored_document_is_untouched_by_a_new_template(
        self, faculty, committee, issue_order, order_template
    ):
        from django.conf import settings
        from django.core.files import File

        from orders.order_template import install_order_template

        order = issue_order(faculty=faculty, committee=committee)
        with order.generated_file.open("rb") as handle:
            before = handle.read()

        source = settings.BASE_DIR / "tests" / "fixtures" / "order_template.docx"
        with source.open("rb") as handle:
            replacement = install_order_template(File(handle, name="order_template.docx"))

        order.refresh_from_db()
        with order.generated_file.open("rb") as handle:
            assert handle.read() == before
        assert order.template_used == order_template != replacement


class TestIssueRules:
    def test_only_the_principal_may_issue(self, faculty, other_faculty, committee, year):
        with pytest.raises(NotPrincipal):
            issue_assignment_order(
                committee=committee,
                appointments=[(faculty, AssignmentRole.PRESIDING_OFFICER)],
                order_date=date(2026, 7, 1),
                issued_by=other_faculty,
            )

    def test_a_future_order_date_is_refused(self, faculty, committee, issue_order):
        with pytest.raises(OrderDateInvalid):
            issue_order(
                faculty=faculty,
                committee=committee,
                order_date=timezone.localdate() + timedelta(days=1),
            )
        assert Order.objects.count() == 0

    def test_a_date_outside_the_academic_year_is_refused(
        self, faculty, committee, issue_order
    ):
        with pytest.raises(OrderDateInvalid):
            issue_order(faculty=faculty, committee=committee, order_date=date(2026, 5, 1))

    def test_preview_writes_nothing(self, principal, faculty, committee, year):
        preview = preview_assignment_order(
            committee=committee,
            appointments=[(faculty, AssignmentRole.PRESIDING_OFFICER)],
            order_date=date(2026, 7, 1),
            issued_by=principal,
        )
        assert preview.expected_order_no.endswith("001")
        assert Order.objects.count() == 0
        assert Assignment.objects.count() == 0
        assert not OrderSequence.objects.exists()

    def test_preview_names_the_holder_who_would_be_superseded(
        self, principal, faculty, other_faculty, committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        preview = preview_assignment_order(
            committee=committee,
            appointments=[(other_faculty, AssignmentRole.PRESIDING_OFFICER)],
            order_date=date(2026, 8, 1),
            issued_by=principal,
        )
        assert preview.previous_holder.holder == faculty

    def test_issuing_is_refused_when_the_holder_changed_since_the_preview(
        self, principal, faculty, other_faculty, third_faculty, committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        stale = Assignment.objects.get(faculty=faculty).pk
        issue_order(faculty=other_faculty, committee=committee, order_date=date(2026, 8, 1))
        before = Order.objects.count()

        with pytest.raises(PreviewOutdated):
            issue_assignment_order(
                committee=committee,
                appointments=[(third_faculty, AssignmentRole.PRESIDING_OFFICER)],
                order_date=date(2026, 8, 2),
                issued_by=principal,
                expected_superseded_id=stale,
            )
        assert Order.objects.count() == before


class TestRelinquishmentOrder:
    def test_it_ends_the_assignment_and_stores_a_document(
        self, principal, faculty, committee, issue_order
    ):
        order = issue_order(faculty=faculty, committee=committee)
        assignment = order.created_assignments.get()
        relinquishment = issue_relinquishment_order(
            assignment=assignment,
            order_date=date(2026, 8, 1),
            issued_by=principal,
            remarks="On deputation.",
        )
        assignment.refresh_from_db()

        assert relinquishment.order_type == OrderType.RELINQUISHMENT
        assert relinquishment.template_used is None
        assert assignment.status == AssignmentStatus.RELINQUISHED
        assert assignment.ending_order == relinquishment
        text = document_text(relinquishment)
        assert faculty.full_name in text and "relieved of the duties" in text


class TestCancellation:
    def test_it_cancels_the_order_and_voids_what_it_created(
        self, principal, faculty, committee, issue_order
    ):
        order = issue_order(faculty=faculty, committee=committee)
        corrigendum = cancel_order(
            order=order, reason="Issued against the wrong file.", issued_by=principal
        )
        order.refresh_from_db()
        assignment = order.created_assignments.get()

        assert order.is_cancelled
        assert order.cancelled_by_order == corrigendum
        assert corrigendum.order_type == OrderType.CORRIGENDUM
        assert assignment.status == AssignmentStatus.CANCELLED

    def test_it_restores_the_holder_the_order_had_displaced(
        self, principal, faculty, other_faculty, committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        superseding = issue_order(
            faculty=other_faculty, committee=committee, order_date=date(2026, 8, 1)
        )
        cancel_order(order=superseding, reason="Wrong person named.", issued_by=principal)

        restored = Assignment.objects.get(faculty=faculty)
        replacement = Assignment.objects.get(faculty=other_faculty)
        assert restored.status == AssignmentStatus.ACTIVE
        assert restored.end_date is None
        assert restored.superseded_by is None
        assert replacement.status == AssignmentStatus.CANCELLED

    def test_a_reason_is_required(self, principal, faculty, committee, issue_order):
        order = issue_order(faculty=faculty, committee=committee)
        with pytest.raises(CancellationReasonRequired):
            cancel_order(order=order, reason="   ", issued_by=principal)
        order.refresh_from_db()
        assert not order.is_cancelled

    def test_cancelling_twice_is_refused(self, principal, faculty, committee, issue_order):
        order = issue_order(faculty=faculty, committee=committee)
        cancel_order(order=order, reason="Issued in error.", issued_by=principal)
        order.refresh_from_db()
        with pytest.raises(OrderAlreadyCancelled):
            cancel_order(order=order, reason="Again.", issued_by=principal)

    def test_a_corrigendum_cannot_itself_be_cancelled(
        self, principal, faculty, committee, issue_order
    ):
        order = issue_order(faculty=faculty, committee=committee)
        corrigendum = cancel_order(
            order=order, reason="Issued in error.", issued_by=principal
        )
        with pytest.raises(OrderNotCancellable):
            cancel_order(order=corrigendum, reason="Undo it.", issued_by=principal)

    def test_the_cancelled_order_keeps_its_document(
        self, principal, faculty, committee, issue_order
    ):
        order = issue_order(faculty=faculty, committee=committee)
        cancel_order(order=order, reason="Issued in error.", issued_by=principal)
        order.refresh_from_db()
        assert order.generated_file.storage.exists(order.generated_file.name)
        assert faculty.full_name in document_text(order)

    def test_cancellation_is_audited(self, principal, faculty, committee, issue_order):
        order = issue_order(faculty=faculty, committee=committee)
        cancel_order(order=order, reason="Issued in error.", issued_by=principal)
        assert AuditLogEntry.objects.filter(
            action="ORDER_CANCELLED", object_id=str(order.pk)
        ).exists()
