"""
Screens and flows: what each page shows, and what a form will and will not write.
"""

import re
from datetime import date, timedelta

import pytest
from django.urls import reverse
from django.utils.html import escape
from django.utils import timezone

from accounts.models import Designation, Role
from assignments.models import Assignment, AssignmentRole, AssignmentStatus
from orders.models import Order

pytestmark = pytest.mark.django_db


def select_block(html, name):
    match = re.search(r'<select[^>]*name="%s".*?</select>' % name, html, re.S)
    assert match, f"no <select name={name}> on the page"
    return match.group(0)


class TestDashboard:
    def test_it_shows_only_your_own_committees(
        self, client, faculty, other_faculty, committee, exam_committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        issue_order(faculty=other_faculty, committee=exam_committee)
        client.force_login(faculty)
        html = client.get(reverse("assignments:dashboard")).content.decode()
        assert escape(committee.name) in html
        assert escape(exam_committee.name) not in html

    def test_a_superseded_committee_moves_to_the_past_list(
        self, client, faculty, other_faculty, committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        issue_order(faculty=other_faculty, committee=committee, order_date=date(2026, 8, 1))
        client.force_login(faculty)
        html = client.get(reverse("assignments:dashboard")).content.decode()
        assert "Past committees" in html
        assert other_faculty.full_name in html  # "Handed to ..."
        assert "pill-history" in html

    def test_someone_with_nothing_sees_the_empty_state(self, client, faculty):
        client.force_login(faculty)
        html = client.get(reverse("assignments:dashboard")).content.decode()
        assert "You currently hold no committees" in html


class TestDirectory:
    def test_it_lists_active_assignments_only(
        self, client, faculty, other_faculty, committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        issue_order(faculty=other_faculty, committee=committee, order_date=date(2026, 8, 1))
        client.force_login(faculty)
        html = client.get(reverse("assignments:directory")).content.decode()
        body = html.split("<tbody>", 1)[1]
        assert other_faculty.full_name in body
        assert faculty.full_name not in body

    def test_the_filter_narrows_by_committee_code(
        self, client, faculty, other_faculty, committee, exam_committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        issue_order(faculty=other_faculty, committee=exam_committee)
        client.force_login(faculty)
        html = client.get(reverse("assignments:directory"), {"q": "EXAM"}).content.decode()
        assert html.split("<tbody>", 1)[1].count('data-label="Committee"') == 1

    def test_a_filter_with_no_matches_explains_itself(self, client, faculty):
        client.force_login(faculty)
        html = client.get(reverse("assignments:directory"), {"q": "zzzz"}).content.decode()
        assert "No active assignments match" in html


class TestWorkload:
    def test_it_sorts_by_committees_held(
        self, client, faculty, other_faculty, committee, exam_committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        issue_order(faculty=faculty, committee=exam_committee, order_date=date(2026, 7, 2))
        issue_order(
            faculty=other_faculty,
            committee=committee,
            role=AssignmentRole.MEMBER,
            order_date=date(2026, 7, 3),
        )
        client.force_login(faculty)
        html = client.get(reverse("accounts:workload")).content.decode()
        assert html.index(faculty.full_name) < html.index(other_faculty.full_name)

    def test_it_never_shows_contact_details(
        self, client, faculty, other_faculty, committee, issue_order
    ):
        other_faculty.mobile = "+919000000001"
        other_faculty.save(update_fields=["mobile"])
        issue_order(faculty=other_faculty, committee=committee)
        client.force_login(faculty)
        html = client.get(reverse("accounts:workload")).content.decode()
        assert "+919000000001" not in html
        assert other_faculty.email not in html


class TestProfile:
    def test_the_mobile_number_can_be_changed(self, client, faculty):
        client.force_login(faculty)
        client.post(reverse("accounts:profile"), {"mobile": "+919876543210"})
        faculty.refresh_from_db()
        assert faculty.mobile == "+919876543210"

    def test_an_invalid_number_is_rejected(self, client, faculty):
        client.force_login(faculty)
        response = client.post(reverse("accounts:profile"), {"mobile": "12345"})
        faculty.refresh_from_db()
        assert "Enter a 10-digit mobile number" in response.content.decode()
        assert faculty.mobile != "12345"

    def test_posting_a_role_or_status_changes_nothing(self, client, faculty, department):
        """The most valuable security test in the project (T2)."""
        client.force_login(faculty)
        before = {
            field: getattr(faculty, field)
            for field in ("role", "is_active", "is_staff", "is_superuser", "designation", "email", "full_name")
        }
        client.post(
            reverse("accounts:profile"),
            {
                "mobile": "+919876543210",
                "role": Role.PRINCIPAL,
                "is_active": "False",
                "is_staff": "on",
                "is_superuser": "on",
                "designation": Designation.PRINCIPAL,
                "email": "attacker@evil.example.com",
                "full_name": "Someone Else",
            },
        )
        faculty.refresh_from_db()
        assert {field: getattr(faculty, field) for field in before} == before


def assign_payload(committee, appointments, *, order_date=date(2026, 7, 1), step="preview", **extra):
    """POST data for the assign form, including the formset's management fields."""
    data = {
        "committee": committee.pk if committee else "",
        "order_date": order_date.isoformat() if order_date else "",
        "references": "",
        "remarks": "",
        "step": step,
        "form-TOTAL_FORMS": str(max(len(appointments), 1)),
        "form-INITIAL_FORMS": "0",
        "form-MIN_NUM_FORMS": "0",
        "form-MAX_NUM_FORMS": "20",
    }
    for index, (holder, role) in enumerate(appointments):
        prefix = "x" if hasattr(holder, "organisation") else "u"
        data[f"form-{index}-person"] = f"{prefix}:{holder.pk}" if holder else ""
        data[f"form-{index}-role"] = role
    data.update(extra)
    return data


class TestAssignFlow:
    def test_the_person_dropdown_offers_an_explicit_empty_choice(
        self, principal_client, faculty, committee
    ):
        """Nobody must be pre-selected: the browser would otherwise pick the first person."""
        html = principal_client.get(reverse("orders:assign")).content.decode()
        block = select_block(html, "form-0-person")
        assert 'value="" selected' in block
        assert "Choose a person" in block

    def test_staff_and_external_members_are_offered_separately(
        self, principal_client, faculty, external_member, committee
    ):
        block = select_block(
            principal_client.get(reverse("orders:assign")).content.decode(), "form-0-person"
        )
        assert "Staff" in block and "From outside the college" in block
        assert external_member.organisation in block

    def test_each_name_shows_the_committee_count(
        self, principal_client, faculty, committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        html = principal_client.get(reverse("orders:assign")).content.decode()
        assert f"{faculty.full_name} · CO · 1 active committee" in html

    def test_an_order_with_nobody_named_is_refused(self, principal_client, committee, year):
        response = principal_client.post(
            reverse("orders:assign"), assign_payload(committee, [])
        )
        assert "Name at least one person" in response.content.decode()
        assert Order.objects.count() == 0

    def test_a_person_without_a_role_is_refused(
        self, principal_client, faculty, committee, year
    ):
        response = principal_client.post(
            reverse("orders:assign"), assign_payload(committee, [(faculty, "")])
        )
        assert "Choose a role for this person." in response.content.decode()
        assert Order.objects.count() == 0

    def test_a_choice_survives_a_validation_error(
        self, principal_client, faculty, committee, year
    ):
        response = principal_client.post(
            reverse("orders:assign"), assign_payload(committee, [(faculty, "")])
        )
        block = select_block(response.content.decode(), "form-0-person")
        assert f'value="u:{faculty.pk}" selected' in block

    def test_the_preview_lists_everyone_and_writes_nothing(
        self, principal_client, faculty, other_faculty, committee, year
    ):
        before = (Order.objects.count(), Assignment.objects.count())
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee,
                [
                    (faculty, AssignmentRole.PRESIDING_OFFICER),
                    (other_faculty, AssignmentRole.MEMBER),
                ],
            ),
        )
        html = response.content.decode()
        assert "Issue order" in html
        assert faculty.full_name in html and other_faculty.full_name in html
        assert (Order.objects.count(), Assignment.objects.count()) == before

    def test_one_order_appoints_several_people(
        self, principal_client, faculty, other_faculty, third_faculty, external_member, committee, year
    ):
        """The college's real orders name between two and eight people."""
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee,
                [
                    (faculty, AssignmentRole.PRESIDING_OFFICER),
                    (other_faculty, AssignmentRole.MEMBER),
                    (third_faculty, AssignmentRole.MEMBER),
                    (external_member, AssignmentRole.EXTERNAL_MEMBER),
                ],
                step="issue",
                expected_superseded="none",
            ),
        )
        order = Order.objects.get()
        assert response.status_code == 302
        assert order.created_assignments.count() == 4
        assert order.created_assignments.filter(external_member=external_member).exists()
        assert order.render_context["member_count"] == 4
        assert [row["name"] for row in order.render_context["member_list"]] == [
            faculty.full_name,
            other_faculty.full_name,
            third_faculty.full_name,
            external_member.full_name,
        ]

    def test_two_people_in_charge_on_one_order_is_refused(
        self, principal_client, faculty, other_faculty, committee, year
    ):
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee,
                [
                    (faculty, AssignmentRole.PRESIDING_OFFICER),
                    (other_faculty, AssignmentRole.NODAL_OFFICER),
                ],
            ),
        )
        assert "only one Presiding Officer or Nodal Officer" in response.content.decode()
        assert Order.objects.count() == 0

    def test_a_clash_names_the_person_and_marks_their_row(
        self, principal_client, faculty, other_faculty, committee, issue_order
    ):
        """With several people on one order, "this member" would not say who."""
        issue_order(faculty=faculty, committee=committee, role=AssignmentRole.MEMBER)
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee,
                [
                    (other_faculty, AssignmentRole.PRESIDING_OFFICER),
                    (faculty, AssignmentRole.MEMBER),
                ],
            ),
        )
        html = response.content.decode()
        assert f"{faculty.full_name} is already Member of {escape(committee.name)}" in html
        formset = response.context["formset"]
        assert formset.forms[1].non_field_errors(), "the error belongs to the second row"
        assert not formset.forms[0].non_field_errors()
        assert Order.objects.count() == 1

    def test_an_external_clash_does_not_call_them_a_faculty_member(
        self, principal_client, external_member, committee, issue_order
    ):
        issue_order(
            committee=committee,
            appointments=[(external_member, AssignmentRole.EXTERNAL_MEMBER)],
        )
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee, [(external_member, AssignmentRole.EXTERNAL_MEMBER)]
            ),
        )
        html = response.content.decode()
        assert f"{external_member.full_name} is already External Member" in html
        assert "faculty member" not in html

    def test_the_same_person_twice_on_one_order_is_refused(
        self, principal_client, faculty, committee, year
    ):
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee,
                [(faculty, AssignmentRole.MEMBER), (faculty, AssignmentRole.MEMBER)],
            ),
        )
        assert "named twice" in response.content.decode()
        assert Order.objects.count() == 0

    def test_an_external_member_cannot_take_a_staff_role(
        self, principal_client, external_member, committee, year
    ):
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(committee, [(external_member, AssignmentRole.PRESIDING_OFFICER)]),
        )
        assert "only be appointed as External Member" in response.content.decode()
        assert Order.objects.count() == 0

    def test_a_staff_member_cannot_take_the_external_role(
        self, principal_client, faculty, committee, year
    ):
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(committee, [(faculty, AssignmentRole.EXTERNAL_MEMBER)]),
        )
        assert "for people outside the college" in response.content.decode()
        assert Order.objects.count() == 0

    def test_issuing_creates_the_order_and_the_assignment(
        self, principal_client, faculty, committee, year
    ):
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee,
                [(faculty, AssignmentRole.PRESIDING_OFFICER)],
                step="issue",
                expected_superseded="none",
            ),
        )
        order = Order.objects.get()
        assert response.status_code == 302
        assert response["Location"] == reverse("orders:detail", kwargs={"pk": order.pk})
        assert order.created_assignments.get().faculty == faculty

    def test_the_preview_describes_people_as_the_order_does(
        self, principal_client, faculty, external_member, committee, year
    ):
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee,
                [
                    (faculty, AssignmentRole.PRESIDING_OFFICER),
                    (external_member, AssignmentRole.EXTERNAL_MEMBER),
                ],
            ),
        )
        html = response.content.decode()
        assert escape(f"{faculty.get_designation_display()}, {faculty.department.name}") in html
        assert escape(external_member.organisation) in html

    def test_people_are_listed_in_the_order_they_were_named(
        self, principal_client, faculty, other_faculty, third_faculty, committee, year
    ):
        principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee,
                [
                    (third_faculty, AssignmentRole.PRESIDING_OFFICER),
                    (faculty, AssignmentRole.MEMBER),
                    (other_faculty, AssignmentRole.MEMBER),
                ],
                step="issue",
                expected_superseded="none",
            ),
        )
        order = Order.objects.get()
        html = principal_client.get(
            reverse("orders:detail", kwargs={"pk": order.pk})
        ).content.decode()
        positions = [
            html.index(person.full_name)
            for person in (third_faculty, faculty, other_faculty)
        ]
        assert positions == sorted(positions), "listed in the order they were named"

    def test_rows_added_beyond_the_first_three_are_accepted(
        self, principal_client, faculty, other_faculty, third_faculty, external_member, committee, year
    ):
        """The "Add another person" button extends the formset past its initial rows."""
        principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee,
                [
                    (faculty, AssignmentRole.PRESIDING_OFFICER),
                    (other_faculty, AssignmentRole.MEMBER),
                    (third_faculty, AssignmentRole.MEMBER),
                    (external_member, AssignmentRole.EXTERNAL_MEMBER),
                ],
                step="issue",
                expected_superseded="none",
            ),
        )
        assert Order.objects.get().created_assignments.count() == 4

    def test_empty_rows_between_filled_ones_are_ignored(
        self, principal_client, faculty, other_faculty, committee, year
    ):
        payload = assign_payload(
            committee,
            [
                (faculty, AssignmentRole.PRESIDING_OFFICER),
                (None, ""),
                (other_faculty, AssignmentRole.MEMBER),
            ],
            step="issue",
            expected_superseded="none",
        )
        principal_client.post(reverse("orders:assign"), payload)
        assert Order.objects.get().created_assignments.count() == 2

    def test_a_committee_without_a_template_cannot_be_used(
        self, principal_client, faculty, db, year
    ):
        from committees.models import Committee

        bare = Committee.objects.create(code="NEW", name="Committee Without Template")
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(bare, [(faculty, AssignmentRole.PRESIDING_OFFICER)]),
        )
        assert "no active order template" in response.content.decode()
        assert Order.objects.count() == 0

    def test_issuing_the_same_preview_twice_does_not_duplicate(
        self, principal_client, faculty, committee, year
    ):
        """A second click of Issue order must not produce a second order."""
        payload = assign_payload(
            committee,
            [(faculty, AssignmentRole.PRESIDING_OFFICER)],
            step="issue",
            expected_superseded="none",
        )
        first = principal_client.post(reverse("orders:assign"), payload)
        second = principal_client.post(reverse("orders:assign"), payload)
        assert first.status_code == 302
        assert second.status_code == 200, "the repeat is refused, not redirected"
        assert "is already Presiding Officer" in second.content.decode()
        assert Order.objects.count() == 1

    def test_issuing_without_a_preview_is_refused(
        self, principal_client, faculty, committee, year
    ):
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee, [(faculty, AssignmentRole.PRESIDING_OFFICER)], step="issue"
            ),
        )
        assert "previewed" in response.content.decode()
        assert Order.objects.count() == 0

    def test_a_future_date_is_refused_with_a_readable_message(
        self, principal_client, faculty, committee, year
    ):
        response = principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee,
                [(faculty, AssignmentRole.PRESIDING_OFFICER)],
                order_date=timezone.localdate() + timedelta(days=1),
            ),
        )
        assert "cannot be in the future" in response.content.decode()

    def test_references_are_stored_and_printed_on_the_order(
        self, principal_client, faculty, committee, year
    ):
        reference = "GTU circular no. 2026-27/456 dated 04-02-2026"
        principal_client.post(
            reverse("orders:assign"),
            assign_payload(
                committee,
                [(faculty, AssignmentRole.PRESIDING_OFFICER)],
                step="issue",
                expected_superseded="none",
                references=reference,
            ),
        )
        order = Order.objects.get()
        assert order.references == reference
        assert order.render_context["references"] == reference


class TestFacultyManagement:
    def test_a_new_record_is_always_created_as_faculty(self, principal_client, department):
        principal_client.post(
            reverse("accounts:faculty_create"),
            {
                "salutation": "Shri",
                "full_name": "Test Faculty Nine",
                "email": "New.Person@test.local",
                "designation": Designation.LECTURER,
                "govt_class": "CLASS_II",
                "department": department.pk,
                "employee_code": "",
                "joining_date": "",
                "mobile": "",
                "role": Role.PRINCIPAL,
                "is_staff": "on",
                "is_superuser": "on",
            },
        )
        from accounts.models import User

        person = User.objects.get(email="new.person@test.local")
        assert person.role == Role.FACULTY
        assert not person.is_staff and not person.is_superuser

    def test_deactivation_does_not_end_committees(
        self, principal_client, faculty, committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        response = principal_client.post(
            reverse("accounts:faculty_deactivate", kwargs={"pk": faculty.pk}),
            {"confirm": "on"},
        )
        faculty.refresh_from_db()
        assert response.status_code == 302
        assert not faculty.is_active
        assert Assignment.objects.active().filter(faculty=faculty).count() == 1

    def test_the_deactivation_page_warns_about_committees_held(
        self, principal_client, faculty, committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        html = principal_client.get(
            reverse("accounts:faculty_deactivate", kwargs={"pk": faculty.pk})
        ).content.decode()
        assert "still holds" in html and escape(committee.name) in html


class TestCommitteeManagement:
    def test_a_lowercase_code_is_corrected_rather_than_refused(self, principal_client):
        principal_client.post(
            reverse("committees:create"),
            {"name": "Test Committee One", "code": "tst", "allows_multiple_members": "on", "is_active": "on"},
        )
        from committees.models import Committee

        assert Committee.objects.get(name="Test Committee One").code == "TST"

    def test_the_code_cannot_change_once_an_order_exists(
        self, principal_client, faculty, committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        principal_client.post(
            reverse("committees:edit", kwargs={"pk": committee.pk}),
            {"name": committee.name, "code": "HACKED", "allows_multiple_members": "on", "is_active": "on"},
        )
        committee.refresh_from_db()
        assert committee.code == "TPU"


class TestTemplateUpload:
    def _upload(self, client, committee, name, content):
        from django.core.files.uploadedfile import SimpleUploadedFile

        return client.post(
            reverse("committees:template", kwargs={"pk": committee.pk}),
            {"docx_file": SimpleUploadedFile(name, content)},
        )

    def test_a_macro_enabled_file_is_rejected(self, principal_client, committee):
        from django.conf import settings

        source = settings.BASE_DIR / "tests" / "fixtures" / "order_template_TPU.docx"
        response = self._upload(principal_client, committee, "t.docm", source.read_bytes())
        assert "Only .docx files are accepted" in response.content.decode()
        assert committee.templates.count() == 1

    def test_a_file_that_is_not_a_document_is_rejected(self, principal_client, committee):
        response = self._upload(principal_client, committee, "t.docx", b"not a zip at all")
        assert "not a valid Word document" in response.content.decode()
        assert committee.templates.count() == 1

    def test_a_template_missing_placeholders_names_them(self, principal_client, committee):
        import io

        from docx import Document

        document = Document()
        document.add_paragraph("Nothing useful here.")
        buffer = io.BytesIO()
        document.save(buffer)
        response = self._upload(principal_client, committee, "bare.docx", buffer.getvalue())
        body = response.content.decode()
        assert "missing" in body and "order_no" in body
        assert committee.templates.count() == 1

    def test_a_valid_upload_supersedes_the_previous_version(
        self, principal_client, committee
    ):
        from django.conf import settings

        source = settings.BASE_DIR / "tests" / "fixtures" / "order_template_EXAM.docx"
        self._upload(principal_client, committee, "v2.docx", source.read_bytes())
        assert committee.templates.count() == 2
        assert committee.templates.filter(is_active=True).count() == 1
        assert committee.templates.order_by("-version").first().version == 2


class TestAuditLogScreen:
    def test_it_shows_readable_action_names(
        self, principal_client, faculty, committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        html = principal_client.get(reverse("audit:log")).content.decode()
        body = html.split("<tbody>", 1)[1]
        assert "Order issued" in body
        assert "ORDER_ISSUED" not in body

    def test_it_can_be_filtered_by_action(
        self, principal_client, faculty, committee, issue_order
    ):
        issue_order(faculty=faculty, committee=committee)
        html = principal_client.get(reverse("audit:log"), {"action": "ORDER_ISSUED"}).content.decode()
        body = html.split("<tbody>", 1)[1]
        assert body.count('data-label="When"') == 1


class TestSignIn:
    def test_the_address_may_be_typed_in_any_case(self, client, faculty):
        faculty.set_password("demo-password")
        faculty.save(update_fields=["password"])
        response = client.post(
            reverse("accounts:login"),
            {"username": "FACULTY.ONE@TEST.LOCAL", "password": "demo-password"},
        )
        assert response.status_code == 302

    def test_a_wrong_password_is_reported_clearly(self, client, faculty):
        response = client.post(
            reverse("accounts:login"),
            {"username": faculty.email, "password": "wrong"},
        )
        assert "not correct" in response.content.decode()

    def test_an_off_site_next_is_ignored(self, client, faculty):
        faculty.set_password("demo-password")
        faculty.save(update_fields=["password"])
        response = client.post(
            reverse("accounts:login"),
            {
                "username": faculty.email,
                "password": "demo-password",
                "next": "https://evil.example.com/",
            },
        )
        assert response["Location"] == reverse("assignments:dashboard")

    def test_a_deactivated_person_cannot_sign_in(self, client, faculty):
        faculty.set_password("demo-password")
        faculty.is_active = False
        faculty.save(update_fields=["password", "is_active"])
        response = client.post(
            reverse("accounts:login"),
            {"username": faculty.email, "password": "demo-password"},
        )
        assert response.status_code == 200
