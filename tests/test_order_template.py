"""The one order template: validation, installation and the setup command."""

import io

import pytest
from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import CommandError, call_command
from docx import Document

from core.exceptions import TemplateRejected
from orders.models import OrderTemplate
from orders.order_template import install_order_template

pytestmark = pytest.mark.django_db
GOOD = settings.BASE_DIR / "tests" / "fixtures" / "order_template.docx"


def docx_bytes(text):
    document = Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_installing_a_template_makes_it_the_only_one_in_use(order_template):
    newer = install_order_template(SimpleUploadedFile("t.docx", GOOD.read_bytes()))
    order_template.refresh_from_db()
    assert newer.is_active and not order_template.is_active
    assert OrderTemplate.objects.filter(is_active=True).count() == 1
    assert OrderTemplate.objects.count() == 2, "the earlier one is kept"


def test_a_macro_enabled_file_is_rejected(order_template):
    with pytest.raises(TemplateRejected, match="Only .docx files are accepted"):
        install_order_template(SimpleUploadedFile("t.docm", GOOD.read_bytes()))
    assert OrderTemplate.objects.active() == order_template


def test_a_file_that_is_not_a_document_is_rejected(order_template):
    with pytest.raises(TemplateRejected, match="not a valid Word document"):
        install_order_template(SimpleUploadedFile("t.docx", b"not a zip at all"))


def test_a_template_missing_placeholders_names_them(order_template):
    with pytest.raises(TemplateRejected) as excinfo:
        install_order_template(SimpleUploadedFile("bare.docx", docx_bytes("Nothing here.")))
    assert "order_no" in str(excinfo.value)
    assert OrderTemplate.objects.active() == order_template


def test_the_setup_command_installs_a_template(order_template):
    call_command("load_order_template", str(GOOD))
    assert OrderTemplate.objects.active() != order_template


def test_the_setup_command_reports_a_bad_file(tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip at all")
    with pytest.raises(CommandError, match="not a valid Word document"):
        call_command("load_order_template", str(bad))


def test_no_page_mentions_template_versions(principal_client, faculty, committee, issue_order):
    order = issue_order(faculty=faculty, committee=committee)
    for url in ("/committees/", f"/committees/{committee.pk}/", f"/orders/{order.pk}/"):
        html = principal_client.get(url).content.decode()
        assert "version" not in html.lower(), url
        assert "Upload order template" not in html, url
