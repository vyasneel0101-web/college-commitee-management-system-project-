"""
Committee order templates: validation and versioned upload. docs/03_SECURITY.md §5.
"""

import zipfile

from django.db import transaction
from django.utils.translation import gettext as _
from docx.opc.exceptions import OpcError
from docxtpl import DocxTemplate
from jinja2 import TemplateError
from jinja2.sandbox import SandboxedEnvironment
from lxml.etree import XMLSyntaxError

from audit import services as audit
from audit.models import AuditAction
from core.exceptions import TemplateRejected

from .models import CommitteeTemplate

REQUIRED_PLACEHOLDERS = frozenset(
    {
        "order_no",
        "order_date",
        "faculty_salutation",
        "faculty_name",
        "faculty_designation",
        "faculty_department",
        "committee_name",
        "assignment_role",
        "academic_year",
        "principal_name",
        "principal_signature",
    }
)
MAX_TEMPLATE_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
ZIP_MAGIC = b"PK\x03\x04"


def validate_template_file(uploaded_file):
    """
    Run the full validation chain and return the sorted placeholder names found.

    Raises TemplateRejected naming the problem. Writes nothing. Order of
    checks: extension, size, magic bytes, zip structure, zip-bomb guard,
    macro parts, parse, required placeholders.
    """
    if not (uploaded_file.name or "").lower().endswith(".docx"):
        raise TemplateRejected(
            _("Only .docx files are accepted. Macro-enabled files such as .docm are rejected.")
        )
    if uploaded_file.size > MAX_TEMPLATE_BYTES:
        raise TemplateRejected(_("The template must be 2 MB or smaller."))

    uploaded_file.seek(0)
    if uploaded_file.read(4) != ZIP_MAGIC:
        raise TemplateRejected(_("The file is not a valid Word document."))
    uploaded_file.seek(0)
    try:
        with zipfile.ZipFile(uploaded_file) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise TemplateRejected(_("The file is not a valid Word document.")) from exc

    if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
        raise TemplateRejected(_("The template expands to more than 20 MB and was rejected."))
    names = {entry.filename.lower() for entry in entries}
    if any(name.endswith("vbaproject.bin") for name in names):
        raise TemplateRejected(
            _("The template contains macros. Save it as a plain .docx and upload again.")
        )
    if "word/document.xml" not in names:
        raise TemplateRejected(_("The file is not a valid Word document."))

    uploaded_file.seek(0)
    try:
        found = DocxTemplate(uploaded_file).get_undeclared_template_variables(
            jinja_env=SandboxedEnvironment()
        )
    except (OpcError, KeyError, ValueError, XMLSyntaxError, TemplateError, zipfile.BadZipFile) as exc:
        raise TemplateRejected(
            _("The template could not be read: %(error)s") % {"error": exc}
        ) from exc
    finally:
        uploaded_file.seek(0)

    missing = sorted(REQUIRED_PLACEHOLDERS - found)
    if missing:
        raise TemplateRejected(
            _("The template is missing %(names)s. Add them and upload again.")
            % {"names": ", ".join("{{ %s }}" % name for name in missing)}
        )
    return sorted(found)


@transaction.atomic
def upload_committee_template(*, committee, uploaded_file, uploaded_by):
    """
    Validate `uploaded_file` and store it as the committee's new active template.

    Invariants: invalid files are rejected before anything is written;
    versions increase by one per committee and are never reused; at most one
    version is active; earlier versions are kept, because past orders were
    rendered from them (BR-7).
    """
    placeholders = validate_template_file(uploaded_file)

    latest = committee.templates.select_for_update().order_by("-version").first()
    template = CommitteeTemplate(
        committee=committee,
        version=latest.version + 1 if latest else 1,
        uploaded_by=uploaded_by,
        detected_placeholders=placeholders,
        is_active=True,
    )
    committee.templates.filter(is_active=True).update(is_active=False)
    template.docx_file.save("template.docx", uploaded_file, save=False)
    try:
        template.save()
    except Exception:
        # The transaction discards the row; the stored file must not outlive it.
        template.docx_file.storage.delete(template.docx_file.name)
        raise

    audit.log(
        action=AuditAction.TEMPLATE_UPLOADED,
        actor=uploaded_by,
        obj=template,
        summary=f"{committee.code} order template version {template.version} uploaded",
        detail={"version": template.version, "placeholders": placeholders},
    )
    return template
